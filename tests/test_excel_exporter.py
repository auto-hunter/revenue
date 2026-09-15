from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from models.schemas import TransactionRecord
from services.excel_exporter import create_excel


def test_creates_downloadable_workbook():
    record = TransactionRecord(
        product_spec="34/0.18TA", gross_weight=367, bobbin_weight=48, net_weight=319, subtotal=1026.5,
        customer="일산전선", corporate_entity="태일전선",
        transaction_date=date(2026, 8, 26), source_file="sample.pdf", page_number=1,
    )
    workbook = load_workbook(filename=BytesIO(create_excel([record])))
    assert workbook.sheetnames == ["추출 결과"]
    assert workbook["추출 결과"]["A2"].value == "34/0.18TA"
    assert workbook["추출 결과"]["I1"].value == "거래법인"
    assert workbook["추출 결과"]["I2"].value == "태일전선"
    for cell_reference in ("B2", "C2", "D2", "E2"):
        assert workbook["추출 결과"][cell_reference].number_format == "0.0"
